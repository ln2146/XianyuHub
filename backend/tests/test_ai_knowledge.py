import io

import pytest

from app.domains.message.knowledge import UnsupportedFile, extract_text, select_from_docs
from app.intelligence.llm import LLMGateway, LLMResult, set_llm
from tests.test_messages import im  # noqa: F401


def _pdf(text: str) -> bytes:
    """手写一个只含一行文字的最小 PDF。"""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def test_extract_text_from_supported_files():
    assert extract_text("faq.txt", "发货：拍下后自动发送".encode("gb18030")) == "发货：拍下后自动发送"
    assert "Shipping within 24h" in extract_text("a.pdf", _pdf("Shipping within 24h"))

    from docx import Document

    doc = Document()
    doc.add_paragraph("退换说明：虚拟商品发出后不退")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "月卡"
    table.rows[0].cells[1].text = "9.9 元"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("说明.docx", buf.getvalue())
    assert "虚拟商品发出后不退" in text and "月卡 | 9.9 元" in text  # 表格里的价目也要读到

    with pytest.raises(UnsupportedFile, match="暂不支持"):
        extract_text("a.xlsx", b"x")
    with pytest.raises(UnsupportedFile, match="没有从文件里读到文字"):
        extract_text("empty.txt", b"   ")


def test_only_relevant_doc_sections_go_into_the_prompt():
    filler = "\n\n".join(f"第{i}章 与发货无关的冗长说明内容" * 5 for i in range(40))
    docs = [("说明书.txt", filler + "\n\n发货时间：每天晚上八点统一发货"), ("价目表.txt", "月卡 9.9 元")]
    picked = select_from_docs(docs, "什么时候发货")
    assert "晚上八点统一发货" in picked and len(picked) <= 1500


class CaptureLLM:
    def __init__(self):
        self.prompts = []

    async def complete(self, **kw):
        if "intent" in kw["schema"]["properties"]:
            return LLMResult(text='{"intent": "logistics"}', model="m")
        self.prompts.append(kw["messages"][-1].content)
        return LLMResult(text='{"reply":"晚上八点统一发哈","need_human":false,"reason":"资料"}', model="m")


async def test_trial_uses_seller_instructions_and_uploaded_docs(admin_client, im):  # noqa: F811
    account_id, _ = im
    await admin_client.patch(
        f"/api/v1/accounts/{account_id}",
        json={"settings": {"ai_instructions": "我是学生党，周末不发货，说话亲切一点"}},
    )
    r = await admin_client.post(
        "/api/v1/messages/knowledge-docs",
        files={"file": ("发货说明.txt", "发货时间：每天晚上八点统一发货".encode(), "text/plain")},
    )
    assert r.status_code == 201 and r.json()["chars"] > 0 and r.json()["account_id"] is None
    bad = await admin_client.post("/api/v1/messages/knowledge-docs", files={"file": ("a.exe", b"x", "x")})
    assert bad.status_code == 400 and "暂不支持" in bad.json()["detail"]

    llm = CaptureLLM()
    set_llm(LLMGateway(provider=llm))
    try:
        # 不需要已有会话
        res = (
            await admin_client.post(
                "/api/v1/messages/ai-trial", json={"account_id": account_id, "text": "什么时候发货"}
            )
        ).json()
    finally:
        set_llm(None)
    assert res["text"] == "晚上八点统一发哈" and res["used_docs"]
    prompt = llm.prompts[-1]
    assert "【店主说明】" in prompt and "周末不发货" in prompt
    assert "【资料】" in prompt and "晚上八点统一发货" in prompt

    # 停用后不再带入
    doc_id = (await admin_client.get("/api/v1/messages/knowledge-docs")).json()[0]["id"]
    await admin_client.patch(f"/api/v1/messages/knowledge-docs/{doc_id}", json={"enabled": False})
    set_llm(LLMGateway(provider=llm))
    try:
        await admin_client.post("/api/v1/messages/ai-trial", json={"account_id": account_id, "text": "什么时候发货"})
    finally:
        set_llm(None)
    assert "【资料】" not in llm.prompts[-1]
    assert (await admin_client.delete(f"/api/v1/messages/knowledge-docs/{doc_id}")).status_code == 204
